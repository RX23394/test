"""Tests for DOCX, UTF-8 text, and Markdown parsers."""

from pathlib import Path

from docx import Document

from app.services.document_parser import (
    DocumentContent,
    DocumentPage,
    DocxParser,
    MarkdownParser,
    TxtParser,
)
from app.services.document_parser.dependencies import get_parser_registry


def test_docx_parser_reads_non_empty_paragraphs(tmp_path: Path) -> None:
    """DOCX paragraphs are returned in order as one logical page."""
    path = tmp_path / "network-guide.docx"
    document = Document()
    document.add_heading("Network Guide", level=1)
    document.add_paragraph("")
    document.add_paragraph("Routing fundamentals")
    document.save(path)

    result = DocxParser().parse(path)

    assert result == DocumentContent(
        filename="network-guide.docx",
        pages=[
            DocumentPage(
                page_number=1,
                text="Network Guide\nRouting fundamentals",
            )
        ],
    )


def test_txt_parser_reads_utf8_text(tmp_path: Path) -> None:
    """Plain-text files are decoded strictly as UTF-8."""
    path = tmp_path / "utf8.txt"
    utf8_text = "\u7f51\u7edc\u914d\u7f6e\nUTF-8 \u5185\u5bb9"
    path.write_text(utf8_text, encoding="utf-8")

    result = TxtParser().parse(path)

    assert result == DocumentContent(
        filename="utf8.txt",
        pages=[DocumentPage(page_number=1, text=utf8_text)],
    )


def test_markdown_parser_preserves_heading_hierarchy(tmp_path: Path) -> None:
    """Markdown ATX heading markers and levels remain unchanged."""
    path = tmp_path / "routing.md"
    markdown = "# Network\n\n## Routing\n\n### OSPF\n\nDetails"
    path.write_text(markdown, encoding="utf-8")

    result = MarkdownParser().parse(path)

    assert result == DocumentContent(
        filename="routing.md",
        pages=[DocumentPage(page_number=1, text=markdown)],
    )


def test_empty_text_documents_have_no_pages(tmp_path: Path) -> None:
    """Empty TXT and Markdown files produce a valid empty page list."""
    txt_path = tmp_path / "empty.txt"
    markdown_path = tmp_path / "empty.markdown"
    txt_path.write_text("  \n", encoding="utf-8")
    markdown_path.write_text("", encoding="utf-8")

    assert TxtParser().parse(txt_path).pages == []
    assert MarkdownParser().parse(markdown_path).pages == []


def test_default_registry_contains_text_document_parsers() -> None:
    """The default registry resolves every newly supported extension."""
    get_parser_registry.cache_clear()
    registry = get_parser_registry()

    assert isinstance(registry.get(Path("document.DOCX")), DocxParser)
    assert isinstance(registry.get(Path("document.TXT")), TxtParser)
    assert isinstance(registry.get(Path("document.MD")), MarkdownParser)
    assert isinstance(registry.get(Path("document.MARKDOWN")), MarkdownParser)
